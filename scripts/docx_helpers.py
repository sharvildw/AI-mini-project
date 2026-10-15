import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def add_callout_box(doc, text, title=None, border_color="3182CE", bg_color="EBF8FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only (academic callout)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="none"/>'
                        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
                        f'<w:bottom w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    if title:
        run_t = p.add_run(title + "\n")
        run_t.font.name = "Times New Roman"
        run_t.font.size = Pt(11)
        run_t.font.bold = True
        run_t.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        
    run_txt = p.add_run(text)
    run_txt.font.name = "Times New Roman"
    run_txt.font.size = Pt(10.5)
    run_txt.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    doc.add_paragraph() # Spacing

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F7FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>'
                        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="CBD5E0"/>'
                        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>'
                        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E0", sz="6", val="single")
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "1A365D") # Deep Navy
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Body rows
    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = "F7FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
                
    if col_widths:
        for row in tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)
    return tbl

def add_figure_image(doc, img_path, caption_text, width_inches=6.0):
    if not os.path.exists(img_path):
        p_err = doc.add_paragraph(f"[Image file not found: {img_path}]")
        p_err.runs[0].font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(10)
    run_cap.font.italic = True
    run_cap.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

print("Helper functions initialized successfully.")
